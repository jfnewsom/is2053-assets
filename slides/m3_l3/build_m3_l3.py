"""
build_m3_l3.py - M3-L3 ("Lab 3.3: Trip Report") walkthrough deck.

Module 3, Lab 3 (MASTER) | walks all seven checkpoints of the current sheet.

Source-of-truth chain (do not invert):
  1. Modules/mod3/lab-3-3.py          - bedrock main game code
  2. Modules/mod3/function_library.py - bedrock library code
  3. pages/labs/json/lab-3-3.json     - student-facing assignment
  4. This deck (downstream)

Checkpoint numbers match the sheet exactly (CP1 Cleanup through CP7
save_trip_report). The June 15 deck merged and renumbered them, so its
"Checkpoint 4" was the sheet's CP5; that was most of what read choppy.
No scope or not-permitted claims: if it's in the book through Chapter 7,
it's allowed.

Deck conventions (Oct 2026):
  - Every slide carries a VIDEO SCRIPT. Scripts are written to be spoken:
    contractions, no em-dashes, Oxford comma, no dates, no grading-weight
    claims, and code narrated per the Code Narration Standard.
  - Code and output panels hold at most 13 lines. Code lines are at most
    79 characters, output lines at most 100. The build stops if not.
  - Demo notes say "CODE ON SCREEN (complete as shown):". Nothing tells
    the instructor to type live.
  - Every output panel is a verbatim slice of a real run. The CP2 check
    is the sheet's list check run on the solution's lists; the rest come
    from the solution run under the CodeGrade fixture (randint gives 1 when its upper bound is
    9 or less, otherwise 15). Full Check inputs: 1, TestPlayer, 2, 1, 1,
    n, 3. Austin round trip: 1, TestPlayer, 3, y, 2, 1, 1, n, 3. The
    sorting demo is a real-random run (seed 31) with three hazards.
  - The build also writes the narration file: one paragraph per slide,
    each starting with "[short pause]", no slide numbers. ElevenLabs
    exports one MP3 per paragraph, so MP3 N is slide N.

USAGE
-----
    python3 slides/m3_l3/build_m3_l3.py

Set IS2053_OUT to choose the output folder (default /home/claude).
"""
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

# ----- Bootstrap: make slides/build_lib.py importable -----
SLIDES_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SLIDES_DIR))

from build_lib import (  # noqa: E402
    DeckBuilder,
    format_concept_notes,
    format_output_notes,
)

# ----- Outputs -----
OUT_DIR = Path(os.environ.get('IS2053_OUT', '/home/claude'))
OUTPUT_PPTX = str(OUT_DIR / 'IS2053_2026-10-22_M3-L3_Deck.pptx')
OUTPUT_TTS = str(OUT_DIR / 'IS2053_2026-10-22_M3-L3_TTS.txt')
PNG_OUT = str(OUT_DIR / 'm3_l3_pngs')
WORK_DIR = '/tmp/build_m3_l3'

# ----- Panel limits -----
MAX_LINES = 13
MAX_CODE_COLS = 79
MAX_OUTPUT_COLS = 100
MAX_STORY_WORDS = 75

TTS = []


def say(script):
    """Record one slide's narration for the TTS file and return it."""
    TTS.append('[short pause] ' + ' '.join(script.split()))
    return script


def check_panel(kind, title, text):
    lines = text.rstrip('\n').split('\n')
    cols = MAX_CODE_COLS if kind == 'code' else MAX_OUTPUT_COLS
    if len(lines) > MAX_LINES:
        raise SystemExit(f'{title}: {len(lines)} lines (max {MAX_LINES})')
    wide = [ln for ln in lines if len(ln) > cols]
    if wide:
        raise SystemExit(f'{title}: line over {cols} chars: {wide[0]!r}')


def code_notes(code, instructor_notes):
    """Demo-slide notes. The code on screen is complete as shown."""
    return (
        f'CODE ON SCREEN (complete as shown):\n\n{code.strip()}\n\n\n'
        f'INSTRUCTOR NOTES\n\n{instructor_notes.strip()}'
    )


def title_notes(deck_id, deck_title, opening_line):
    return (
        f'TITLE SLIDE: {deck_id}: {deck_title}\n\n'
        'This is the cold open. Brief greeting, name the unit, then '
        'advance.\n\n'
        f'Suggested opening line:\n{opening_line}'
    )


def demo(deck, title, code, png, script, instructor, key_terms=None,
         think=None):
    check_panel('code', title, code)
    notes = (format_concept_notes(say(script), key_terms, think)
             + '\n\n\n' + code_notes(code, instructor))
    deck.add_demo_slide(title=title, code=code, png=png, notes=notes)


def output(deck, title, text, png, script, instructor, key_terms=None,
           think=None):
    check_panel('output', title, text)
    notes = (format_concept_notes(say(script), key_terms, think)
             + '\n\n\n' + format_output_notes(text, instructor))
    deck.add_output_slide(title=title, output_text=text, png=png,
                          notes=notes)


def concept(deck, title, bullets, script, key_terms=None, think=None,
            source_url=None):
    if not 3 <= len(bullets) <= 5:
        raise SystemExit(f'{title}: {len(bullets)} key points')
    deck.add_concept_slide(
        title=title, bullets=bullets,
        notes=format_concept_notes(say(script), key_terms, think,
                                   source_url))


# ----- Post-processing (same fixes as the Module 4 pipeline) -----
GREEN = RGBColor(0x39, 0xFF, 0x14)
MONO = 'Roboto Mono'


def _restyle_backticks(p):
    """Render `code` spans in green Roboto Mono and drop the backticks."""
    full = ''.join(r.text for r in p.runs)
    if '`' not in full:
        return
    base = p.runs[0].font
    name, size, bold = base.name, base.size, base.bold
    try:
        color = base.color.rgb
    except Exception:
        color = None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for i, seg in enumerate(full.split('`')):
        if seg == '':
            continue
        run = p.add_run()
        run.text = seg
        f = run.font
        if size is not None:
            f.size = size
        if bold is not None:
            f.bold = bold
        if i % 2 == 1:
            f.name = MONO
            f.color.rgb = GREEN
        else:
            if name:
                f.name = name
            if color is not None:
                f.color.rgb = color


def _fix_narrative(slide):
    """Template v2 overlaps the story one-liner and body. Separate them."""
    labels = [sh.text_frame.text.strip() for sh in slide.shapes
              if sh.has_text_frame]
    if 'THE STORY' not in labels:
        return
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        top_in = sh.top / 914400 if sh.top is not None else 0
        if 2.5 < top_in < 2.95:
            sh.height = Emu(int(0.95 * 914400))
        elif 3.1 < top_in < 3.5:
            sh.top = Emu(int(3.78 * 914400))
            sh.height = Emu(int(2.95 * 914400))


def postprocess(path):
    prs = Presentation(path)
    for s in prs.slides:
        _fix_narrative(s)
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if 'Spring 2026' in r.text:
                        r.text = r.text.split('Spring 2026')[0].rstrip(
                            ' •')
                _restyle_backticks(p)
    prs.save(path)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    deck = DeckBuilder(png_out=PNG_OUT, work=WORK_DIR)

    # =========================================================================
    # SLIDE 1: Title
    # =========================================================================
    deck.add_title_slide(
        main_title='Lab 3.3: Trip Report',
        subtitle='IS2053 Programming I  •  Module 3  •  Lab Walkthrough',
        notes=title_notes('M3-L3', 'Lab 3.3: Trip Report', say(
            "Welcome to your Lab 3.3 walkthrough. This is the MASTER lab for "
            "Module 3, and it's the one where the player finally gets "
            "something to keep: a trip report, on screen and saved to a file. "
            "Seven checkpoints, and most of them are short. Let's go."
        )),
    )

    # =========================================================================
    # SLIDE 2: The story
    # =========================================================================
    story_body = (
        "You already have save and load, parallel lists for destinations, and "
        "lists that track the route and the hazards. This lab moves the "
        "hazards into three parallel lists, picks one with a new library "
        "function, sorts the hazards for display, counts the unique cities, "
        "and writes a trip report to trip_report.txt at the end of every game."
    )
    if len(story_body.split()) > MAX_STORY_WORDS:
        raise SystemExit('Story body is over the word limit')
    deck.add_narrative_slide(
        title='The Player Gets a Souvenir',
        section_label='THE STORY',
        bold_oneliner='Two labs of machinery. This one ships the deliverable.',
        body=story_body,
        notes=format_concept_notes(
            say(
                "Here's the plan. Checkpoint 1 is a clean copy of Lab 3.2. "
                "Checkpoints 2 through 4 turn the hazards into data: three "
                "parallel lists, one library function that picks from them, "
                "and a hazard block that shrinks to a single call. Checkpoint "
                "5 sorts the hazards for display without touching the "
                "original list. Checkpoint 6 counts how many different cities "
                "you visited. And Checkpoint 7 is the payoff: a new on-screen "
                "FINAL TRIP REPORT, plus the same report written to a file "
                "the player can keep. The library grows by exactly one "
                "function. Everything else lives in your lab file."
            ),
            key_terms=[
                ('Parallel lists', 'Separate lists where the same index '
                 'describes the same thing.'),
                ('Deliverable', 'Something the program hands the user to '
                 'keep, here trip_report.txt.'),
            ],
            think_about=[
                "Which part of your Lab 3.2 hazard code would you have to "
                "touch to add a seventh hazard?",
            ],
        ),
    )

    # =========================================================================
    # SLIDE 3: CP1 Cleanup
    # =========================================================================
    concept(
        deck, 'Checkpoint 1: Cleanup',
        [
            "Copy `lab-3-2.py` to `lab-3-3.py`. Leave `function_library.py` "
            "where it is.",
            "Update the docstring: Lab 3.3: Trip Report, a new description, "
            "and the chapters.",
            "Play one game to the beach and confirm it matches Lab 3.2.",
            "Check for YOUR ROUTE and HAZARDS ENCOUNTERED before moving on.",
        ],
        "Checkpoint 1 is housekeeping, and it's worth doing carefully. Copy "
        "your Lab 3.2 file to lab dash three dash three dot p-y, and leave "
        "the function library right where it is. You'll add one function to "
        "it in Checkpoint 3. Update the docstring's title, description, and "
        "chapter list. Then play one full game to the beach and make sure "
        "the trip summary still shows YOUR ROUTE and HAZARDS ENCOUNTERED. If "
        "anything's off, fix it now. That way, if something breaks later "
        "today, you know it's something you added today.",
        think=["Why not just keep editing lab-3-2.py?"],
    )

    # =========================================================================
    # SLIDE 4: CP2 concept
    # =========================================================================
    concept(
        deck, 'Checkpoint 2: Build Parallel Hazard Lists',
        [
            "Three lists: `HAZARD_NAMES`, `HAZARD_HEADS_MSGS`, "
            "`HAZARD_TAILS_MSGS`.",
            "Six items each, and the same index describes the same hazard.",
            "They go after the destination lists, before `get_city_data`.",
            "The old if/elif hazard block stays for now. Checkpoint 4 replaces "
            "it.",
        ],
        "You did this for destinations in Lab 3.2. Now it's the hazards' "
        "turn. Three lists: the hazard names, the messages for escaping on "
        "heads, and the messages for getting delayed on tails. Six items in "
        "each, and position zero in all three describes the same hazard. "
        "That's the whole idea of parallel lists: read them straight down by "
        "index, not across. Copy the eighteen strings from the sheet exactly, "
        "in order. Nothing uses the lists yet. Your old hazard block keeps "
        "running until Checkpoint 4.",
        key_terms=[
            ('Index', "An item's position in a list, starting at zero."),
        ],
        think=[
            "What goes wrong if HAZARD_HEADS_MSGS is in a different order "
            "than HAZARD_NAMES?",
        ],
    )

    # =========================================================================
    # SLIDE 5: CP2 demo
    # =========================================================================
    demo(
        deck, 'Demo: Six Hazards by Name',
        r'''# HAZARD DATA AS PARALLEL LISTS (NEW in 3.3)
HAZARD_NAMES = [
    "Buc-ee's appears!",
    "Armadillo crossing!",
    "Speed trap!",
    "Check engine light!",
    "Whataburger cravings!",
    "Toll road confusion!"
]
# HAZARD_HEADS_MSGS and HAZARD_TAILS_MSGS have the same shape:
# six strings each, in the same order as the names
''',
        'slide05_hazard_names.png',
        "Here's the first list. Six hazard names, one per line, with a comma "
        "after every item but the last. Buc-ee's uses double quotes because "
        "of the apostrophe. The heads messages and the tails messages are "
        "built exactly the same way, and the sheet lists all twelve. The "
        "only rule that matters is order: the escape message for Buc-ee's "
        "has to be first in the heads list, because Buc-ee's is first here.",
        "The other two lists are on the sheet (CP2 steps 3 and 4) and in the "
        "next slide's check output. A missing comma between two strings "
        "silently joins them, so the count check catches it.",
    )

    # =========================================================================
    # SLIDE 6: CP2 output
    # =========================================================================
    output(
        deck, 'Output: Lengths Match, Rows Line Up',
        r'''DEBUG: HAZARD_NAMES len      = 6
DEBUG: HAZARD_HEADS_MSGS len = 6
DEBUG: HAZARD_TAILS_MSGS len = 6
DEBUG: All match length 6?   = True
  [0] Buc-ee's appears! | H: You resist the beaver's call. | T: Lost 2 hours on beef jerky.
  [1] Armadillo crossing! | H: It scurries across safely. | T: You swerved. Armadillo unbothered.
  [2] Speed trap! | H: You saw the cop in time. | T: Small town officer says howdy.
''',
        'slide06_list_check_output.png',
        "This is the quick check from the sheet. All three lists report six "
        "items, and the match line says True. Then each row prints one index "
        "across all three lists, so you can read it like a table: Buc-ee's "
        "pairs with the beaver message and the beef jerky delay. If a row "
        "doesn't make sense, your lists are out of order. Once all six rows "
        "look right, delete the check before Checkpoint 3.",
        "First seven lines of the sheet's CP2 check, run on the solution's "
        "lists. Rows 3 to 5 follow; row 4 is 102 characters, so the slice "
        "stops at row 2.",
    )

    # =========================================================================
    # SLIDE 7: CP3 concept
    # =========================================================================
    concept(
        deck, 'Checkpoint 3: Add get_random_hazard() to the Library',
        [
            "One new function in `function_library.py`, after "
            "`get_redirect`.",
            "It takes the three lists as parameters, never your lab's "
            "globals.",
            "It picks one random index and hands back three values.",
            "It's in the library because Module 4 reuses it.",
        ],
        "Checkpoint 3 adds the only new library function this lab. It goes "
        "in a new HAZARD HELPER section, right after get redirect. It takes "
        "the three lists as parameters. The library can't see the lists in "
        "your lab file, and it shouldn't. It picks one random index and hands "
        "back the name and both messages from that same index. Why the "
        "library? Module 4 uses it again when the hazards move into a text "
        "file. A function with a known future use starts in the library.",
        key_terms=[
            ('Tuple return', 'return a, b, c hands back three values at '
             'once; the caller unpacks them into three names.'),
        ],
        think=[
            "Why randint(0, len(names) - 1) and not randint(0, len(names))?",
        ],
    )

    # =========================================================================
    # SLIDE 8: CP3 demo
    # =========================================================================
    demo(
        deck, 'Demo: get_random_hazard()',
        r'''def get_random_hazard(names, heads_msgs, tails_msgs):
    """Pick a random hazard from three parallel lists."""
    index = random.randint(0, len(names) - 1)
    return names[index], heads_msgs[index], tails_msgs[index]
''',
        'slide08_get_random_hazard.png',
        "Two lines do all the work. We pick a random index from zero up to "
        "one less than the length of the names list, because the last valid "
        "index is always one less than the length. Then we hand back three "
        "values: the name, the heads message, and the tails message, all at "
        "that one index. One index for all three is what keeps them matched. "
        "Give it the full docstring from the sheet, with Args and Returns.",
        "Solution body, lines 332 and 333, with a one-line docstring. The "
        "sheet's skeleton has the full Args/Returns docstring. Test it with "
        "the sheet's two-item lists in the library's main guard.",
    )

    # =========================================================================
    # SLIDE 9: CP4 concept
    # =========================================================================
    concept(
        deck, 'Checkpoint 4: Refactor the In-Game Hazard Block',
        [
            "Out: `random.randint(1, 3)` and the if/elif chain that picked a "
            "hazard.",
            "In: one call that unpacks into `hazard`, `heads_msg`, and "
            "`tails_msg`.",
            "Everything after the call runs once, whichever hazard fired.",
            "All six hazards can now show up in play.",
        ],
        "Now the lists earn their keep. Find the hazard block, everything "
        "under the check percentage if. Whatever shape yours has, replace "
        "everything under that if with the block on the next two slides. The "
        "random number and the if elif chain go away. One call to get random "
        "hazard hands back three values, and we unpack them into three names. "
        "After that, the block does four things once: print the hazard, "
        "append it to the hazards list, flip the coin, and on heads, count "
        "an escape.",
        think=[
            "How many lines would a seventh hazard take now? How many "
            "before?",
        ],
    )

    # =========================================================================
    # SLIDES 10-11: CP4 demo
    # =========================================================================
    demo(
        deck, 'Demo: The Hazard Block, Part 1',
        r'''if function_library.check_percentage(HAZARD_CHANCE):
    hazards_faced = hazards_faced + 1
    print()
    print('*** HAZARD! ***')

    hazard, heads_msg, tails_msg = (
        function_library.get_random_hazard(
            HAZARD_NAMES,
            HAZARD_HEADS_MSGS,
            HAZARD_TAILS_MSGS))
    print(hazard)
    hazards_list.append(hazard)
''',
        'slide10_hazard_block_1.png',
        "Here's the top half, exactly as the sheet has it. If the percentage "
        "check comes back True, we add one to hazards faced and print the "
        "hazard banner. Then the call. We pass the three lists to get random "
        "hazard, and the three values that come back are assigned to "
        "hazard, heads message, and tails message, in that order. The "
        "parentheses around the call just let it span several lines. Then we "
        "print the hazard and append it to the hazards list.",
        "Matches the sheet's CP4 block (and the solution) line for line. The "
        "wrapping parentheses keep the unpacking line under 79 characters at "
        "the solution's indent depth.",
    )
    demo(
        deck, 'Demo: The Hazard Block, Part 2',
        r'''    # hazard block, continued: one coin flip for any hazard
    flip = function_library.flip_coin()
    print(f'{flip}!')
    if flip == 'HEADS':
        print(heads_msg)
        hazards_escaped = hazards_escaped + 1
    else:
        print(tails_msg)
''',
        'slide11_hazard_block_2.png',
        "And the bottom half. One coin flip, no matter which hazard fired. "
        "On heads, we print that hazard's escape message and add one to "
        "hazards escaped. On tails, we print its delay message. Before, "
        "these lines were repeated in every branch. Now they're written "
        "once.",
        "Continues the block from the previous slide at the same indent.",
    )

    # =========================================================================
    # SLIDE 12: CP4 output
    # =========================================================================
    output(
        deck, 'Output: A Hazard From the List',
        r'''--- Turn 1 ---
Miles remaining: 200
You rolled a 15!
You travel 150 miles...
Miles remaining: 50

*** HAZARD! ***
Armadillo crossing!
HEADS!
It scurries across safely.
''',
        'slide12_hazard_output.png',
        "Here it is in a run. The hazard check fired, the function picked "
        "Armadillo crossing, the coin came up heads, and the matching escape "
        "message printed. Play a few games and you'll start seeing the new "
        "ones too: check engine light, Whataburger cravings, and toll road "
        "confusion. If you only ever see the first three, look at the range "
        "in get random hazard.",
        "Full Check run under the CodeGrade fixture: randint(0, 5) returns 1, "
        "so the fixture always picks Armadillo crossing!",
    )

    # =========================================================================
    # SLIDE 13: CP5 concept
    # =========================================================================
    concept(
        deck, 'Checkpoint 5: Sort Hazards Alphabetically',
        [
            "Copy first: `sorted_hazards = hazards_list.copy()`.",
            "Then sort the copy: `sorted_hazards.sort()`.",
            "Display `sorted_hazards`; the title becomes HAZARDS ENCOUNTERED "
            "(A-Z).",
            "`hazards_list` stays in arrival order for the save file.",
        ],
        "The sort method changes a list in place. If you sort the hazards "
        "list itself, its arrival order is gone for good, and the save file "
        "needs that order. So we copy first, sort the copy, and display the "
        "copy. It's the same copy pattern you saw in get redirect. The "
        "section title changes too, to HAZARDS ENCOUNTERED with A to Z in "
        "parentheses, in both places it appears.",
        key_terms=[
            ('In place', 'The method changes the list itself and returns '
             'None. hazards_list = hazards_list.sort() leaves you with '
             'None.'),
        ],
        think=["Where does each version of the list get used?"],
    )

    # =========================================================================
    # SLIDE 14: CP5 demo
    # =========================================================================
    demo(
        deck, 'Demo: Copy, Then Sort',
        r'''sorted_hazards = hazards_list.copy()
sorted_hazards.sort()

if len(sorted_hazards) > 0:
    function_library.display_numbered_list(
        sorted_hazards, 'HAZARDS ENCOUNTERED (A-Z)')
else:
    print('HAZARDS ENCOUNTERED (A-Z)')
    print('-' * len('HAZARDS ENCOUNTERED (A-Z)'))
    print('  None - Lucky traveler!')
    print()
''',
        'slide14_copy_sort.png',
        "Two lines make the sorted copy: we copy the hazards list into "
        "sorted hazards, then sort that copy. The display uses sorted "
        "hazards, with the new title. If there were no hazards at all, the "
        "else branch prints the same title by hand, a dash underline sized "
        "to match, and None, Lucky traveler. Because the underline is built "
        "from the length of the title, it grows on its own.",
        "Solution lines 404 to 405 and 428 to 435. The save_game call (if it "
        "fires) still gets hazards_list.",
    )

    # =========================================================================
    # SLIDE 15: CP5 output
    # =========================================================================
    output(
        deck, 'Output: Arrival Order In, A-Z Out',
        r'''HAZARDS ENCOUNTERED (A-Z)
-------------------------
  1. Armadillo crossing!
  2. Buc-ee's appears!
  3. Check engine light!

STATISTICS
----------
  Total Miles: 440
  Total Turns: 6
  Hazards Faced: 3
  Hazards Escaped: 2
  Escape Rate: 66.7%
''',
        'slide15_sorted_output.png',
        "In this game, the hazards actually happened in this order: check "
        "engine light, then Buc-ee's, then the armadillo. The report lists "
        "them alphabetically. The hazards list itself still holds them in "
        "the order they happened, which is what a save file would get.",
        "Real-random solution run (seed 31). Arrival order was Check engine "
        "light!, Buc-ee's appears!, Armadillo crossing!",
    )

    # =========================================================================
    # SLIDE 16: CP6 concept
    # =========================================================================
    concept(
        deck, 'Checkpoint 6: Count Unique Cities Visited',
        [
            "`visited_cities` keeps every arrival, repeats included.",
            "`unique_cities` starts empty and takes each city once, using "
            "`not in`.",
            "Build it before YOUR ROUTE prints, or the count line crashes.",
            "Print the count right after the route.",
        ],
        "The route list is right to show every stop, even when you pass "
        "through San Antonio twice. But a count of different cities "
        "shouldn't double-count. So we build a second list. Start it empty, "
        "loop through the route, and append a city only if it isn't already "
        "in the new list. Build it before the route prints, because the "
        "count line needs it.",
        key_terms=[
            ('Membership test', 'city not in unique_cities is True when the '
             'list does not contain that city yet.'),
        ],
        think=["Go to Austin and back. What should the count be?"],
    )

    # =========================================================================
    # SLIDE 17: CP6 demo
    # =========================================================================
    demo(
        deck, 'Demo: Each City Once',
        r'''unique_cities = []
for city in visited_cities:
    if city not in unique_cities:
        unique_cities.append(city)

# in the report, right after the route:
function_library.display_numbered_list(visited_cities, 'YOUR ROUTE')
print(f'Unique cities visited: {len(unique_cities)}')
print()
''',
        'slide17_unique_cities.png',
        "We set unique cities to an empty list. Then, for each city in the "
        "route, if it isn't in unique cities yet, we append it. The second "
        "time San Antonio comes up, it's already there, so it's skipped. "
        "Later, in the report, the route prints in full, and right under it, "
        "the count of unique cities.",
        "Solution lines 408 to 411 and 423 to 426. The solution splits the "
        "display call over two lines.",
    )

    # =========================================================================
    # SLIDE 18: CP6 output
    # =========================================================================
    output(
        deck, 'Output: Five Stops, Four Cities',
        r'''YOUR ROUTE
----------
  1. San Antonio
  2. Austin
  3. San Antonio
  4. Houston
  5. Galveston

Unique cities visited: 4
''',
        'slide18_unique_output.png',
        "This is the Austin round trip from the sheet: San Antonio to Austin, "
        "back to San Antonio, then Houston and Galveston. Five stops in the "
        "route, four different cities. If your count says five, the not in "
        "check isn't stopping the repeat. If it says three, you're dropping "
        "San Antonio entirely.",
        "Fixture run with inputs 1, TestPlayer, 3, y, 2, 1, 1, n, 3. Same "
        "block as the sheet's CP6 Step 2.",
    )

    # =========================================================================
    # SLIDE 19: CP7 concept
    # =========================================================================
    concept(
        deck, 'Checkpoint 7: save_trip_report() and the Report',
        [
            "`REPORT_FILE = 'trip_report.txt'` joins the constants.",
            "`save_trip_report` lives in `lab-3-3.py`, not the library.",
            "Set `status` first; reset `quit_without_saving = False` every "
            "game.",
            "The on-screen TRIP SUMMARY becomes the FINAL TRIP REPORT.",
            "The file saves after every game. No y/n prompt.",
        ],
        "This is the deliverable. Three pieces. First, a status for the "
        "game: BEACH WINNER, Quit without saving, or Trip Incomplete. For "
        "that, the quit without saving flag has to reset to False at the "
        "top of every game, next to reached beach and saved and quit. "
        "Second, the on-screen summary is replaced by a new layout, the "
        "FINAL TRIP REPORT, and it has to match the sheet line for line. "
        "Third, a new function, save trip report, writes the same "
        "information to trip report dot t-x-t. It stays in your lab file, "
        "because its format belongs to this game, not to the library. And "
        "it saves every game, automatically.",
        key_terms=[
            ("'w' mode", 'Opens a file for writing and replaces whatever was '
             'in it. Every game overwrites the last report.'),
            ('join', "' -> '.join(visited_cities) builds one string with "
             'an arrow between each city.'),
        ],
        think=[
            "Why does save_trip_report stay in the lab file when "
            "get_random_hazard went to the library?",
        ],
    )

    # =========================================================================
    # SLIDE 20: CP7 demo - status, then save
    # =========================================================================
    demo(
        deck, 'Demo: Status First, Save Last',
        r'''if reached_beach:
    status = 'BEACH WINNER!'
elif quit_without_saving:
    status = 'Quit without saving'
else:
    status = 'Trip Incomplete'

# sorted_hazards, unique_cities, and the on-screen report go here
save_trip_report(REPORT_FILE, player_name, current_city,
                 total_miles, total_turns, total_hazards, total_escaped,
                 visited_cities, unique_cities, sorted_hazards, status)
print(f'Trip report saved to {REPORT_FILE}.')
print()
''',
        'slide20_status_and_save.png',
        "After the in-game loop, the first job is the status. If reached "
        "beach is True, status is BEACH WINNER. If the player quit without "
        "saving, that's the status. Otherwise, Trip Incomplete. Then come "
        "the sorted hazards, the unique cities, and the on-screen report "
        "you've already seen. Last, before the play-again prompt, we call "
        "save trip report with all eleven arguments, in order, and print "
        "the confirmation line.",
        "Solution lines 395 to 400 and 452 to 458. The solution wraps the "
        "save_trip_report arguments differently; the order is what matters: "
        "filename, player_name, current_city, total_miles, total_turns, "
        "total_hazards, total_escaped, visited_cities, unique_cities, "
        "sorted_hazards, status.",
    )

    # =========================================================================
    # SLIDES 21-23: CP7 demo - save_trip_report
    # =========================================================================
    demo(
        deck, 'Demo: save_trip_report(), Part 1',
        r'''def save_trip_report(filename, player_name, current_city, total_miles,
                     total_turns, total_hazards, total_escaped,
                     visited_cities, unique_cities, sorted_hazards,
                     status):
    """Write a formatted trip report to disk."""
    outfile = open(filename, 'w')
    outfile.write('DEEP IN THE HEART - TRIP REPORT\n')
    outfile.write('=' * 40 + '\n')
    outfile.write(f'Driver: {player_name}\n')
    outfile.write(f'Final City: {current_city}\n')
    outfile.write(f'Status: {status}\n')
    outfile.write('\n')
''',
        'slide21_report_1.png',
        "Eleven parameters, in the same order as the call. We open the file "
        "in w mode, so each game replaces the last report. Then the header: "
        "the title line, a divider of forty equals signs, and the driver, "
        "final city, and status, each ending with a newline, because write "
        "doesn't add one for you. A blank line closes the section.",
        "Solution lines 103 to 131 with a one-line docstring. The sheet's "
        "skeleton docstring explains why it stays in the lab file.",
    )
    demo(
        deck, 'Demo: save_trip_report(), Part 2',
        r'''    # save_trip_report, continued: route and statistics
    outfile.write('ROUTE:\n')
    route_str = ' -> '.join(visited_cities)
    outfile.write(f'  {route_str}\n')
    outfile.write(f'  Unique cities visited: {len(unique_cities)}\n')
    outfile.write('\n')
    outfile.write('STATISTICS:\n')
    outfile.write(f'  Total Miles: {total_miles}\n')
    outfile.write(f'  Total Turns: {total_turns}\n')
    outfile.write(f'  Hazards Faced: {total_hazards}\n')
    outfile.write(f'  Hazards Escaped: {total_escaped}\n')
''',
        'slide22_report_2.png',
        "The route section joins the visited cities into one string with an "
        "arrow between each pair, and writes it indented two spaces. Then "
        "the unique count and a blank line. The statistics section writes "
        "four totals, each indented two spaces. Notice these are the game "
        "totals, total hazards and total escaped, not the per-leg counters.",
        "Solution lines 133 to 143.",
    )
    demo(
        deck, 'Demo: save_trip_report(), Part 3',
        r'''    if total_hazards > 0:
        rate = (total_escaped / total_hazards) * 100
        outfile.write(f'  Escape Rate: {rate:.1f}%\n')
    outfile.write('\n')
    outfile.write('HAZARDS (A-Z):\n')
    if len(sorted_hazards) > 0:
        for hazard in sorted_hazards:
            outfile.write('  - ' + hazard + '\n')
    else:
        outfile.write('  None - Lucky traveler!\n')
    outfile.write('\n')
    outfile.write('=' * 40 + '\n')
    outfile.close()
''',
        'slide23_report_3.png',
        "The escape rate only makes sense if there was at least one hazard, "
        "so the guard comes first. Without it, a game with no hazards "
        "divides by zero. The rate is written with one decimal place. Then "
        "the hazards section: each sorted hazard on its own line with a "
        "dash, or None, Lucky traveler, if there weren't any. One more "
        "blank line, the closing divider, and we close the file.",
        "The rate line uses the sheet skeleton's f-string format, "
        "{rate:.1f}. The solution builds rate_str with format(rate, '.1f'); "
        "the output is identical. Solution lines 144 to 159.",
    )

    # =========================================================================
    # SLIDE 24: CP7 output - on screen
    # =========================================================================
    output(
        deck, 'Output: The FINAL TRIP REPORT, On Screen',
        r'''==================================================
FINAL TRIP REPORT
==================================================
Driver: TestPlayer
Final City: Galveston
Status: BEACH WINNER!

YOUR ROUTE
----------
  1. San Antonio
  2. Houston
  3. Galveston

''',
        'slide24_final_report_screen.png',
        "Here's the top of the new on-screen report. Equals-sign dividers "
        "top and bottom, then driver, final city, and status first. After "
        "the route come the unique count, the sorted hazards, and the "
        "STATISTICS block, each line indented two spaces, with the escape "
        "rate only when there was a hazard. The sheet shows all twenty-eight "
        "lines. Match them exactly, then the confirmation: Trip report saved "
        "to trip report dot t-x-t.",
        "Full Check run. The rest of the block (unique count, HAZARDS "
        "ENCOUNTERED (A-Z), STATISTICS, closing divider) matches the slices "
        "on the CP5 and CP6 output slides and the sheet's CP7 block.",
    )

    # =========================================================================
    # SLIDES 25-26: CP7 output - the file
    # =========================================================================
    output(
        deck, 'Output: trip_report.txt, Top',
        r'''DEEP IN THE HEART - TRIP REPORT
========================================
Driver: TestPlayer
Final City: Galveston
Status: BEACH WINNER!

ROUTE:
  San Antonio -> Houston -> Galveston
  Unique cities visited: 3

''',
        'slide25_report_file_1.png',
        "And here's the file the player keeps. Open trip report dot t-x-t "
        "after a game. The header, then the route as one line with arrows, "
        "and the unique count.",
        "trip_report.txt written by the Full Check run; matches the sheet's "
        "CP7 Step 2 block.",
    )
    output(
        deck, 'Output: trip_report.txt, Bottom',
        r'''STATISTICS:
  Total Miles: 450
  Total Turns: 3
  Hazards Faced: 1
  Hazards Escaped: 1
  Escape Rate: 100.0%

HAZARDS (A-Z):
  - Armadillo crossing!

========================================
''',
        'slide26_report_file_2.png',
        "The bottom half: the statistics with the escape rate, the hazards "
        "in alphabetical order, and the closing divider. Play two games in "
        "a row and open the file again. It holds only the second game, "
        "because w mode starts the file over every time.",
        "Rest of the same trip_report.txt.",
    )

    # =========================================================================
    # SLIDE 27: Final Check and Submit
    # =========================================================================
    concept(
        deck, 'Final Check and Submit',
        [
            "Full Check inputs: `1`, `TestPlayer`, `2`, `1`, `1`, `n`, `3`.",
            "Match the on-screen FINAL TRIP REPORT line for line.",
            "Check `trip_report.txt`: five sections, overwritten each game.",
            "Try a no-hazard game and an Austin round trip.",
            "Submit `lab-3-3.py` and `function_library.py`. Unlimited "
            "submissions.",
        ],
        "Before you submit, run the Full Check: one for a new game, "
        "TestPlayer, two for Houston, one to continue, one for Galveston, n "
        "to stop playing, and three to quit. Your rolls will differ from the "
        "sheet's, because CodeGrade runs with fixed dice, so compare the "
        "structure line for line. Open trip report dot t-x-t and check all "
        "five sections. Then try the edge cases: a game with no hazards, so "
        "you see None, Lucky traveler, and the Austin round trip, so the "
        "unique count reads four. Submit both files, your lab three three "
        "file and the function library. Submissions are unlimited, so submit "
        "early. That's Module 3. Next up, Module 4 moves the map and the "
        "hazards out of your code and into text files.",
        think=[
            "Which edge case would you test first, and why?",
        ],
    )

    deck.save(OUTPUT_PPTX)
    postprocess(OUTPUT_PPTX)
    if len(TTS) != len(deck.slides):
        raise SystemExit(f'{len(TTS)} narration paragraphs for '
                         f'{len(deck.slides)} slides')
    Path(OUTPUT_TTS).write_text('\n\n'.join(TTS) + '\n', encoding='utf-8')
    print(f'OK {OUTPUT_PPTX}: {len(deck.slides)} slides; '
          f'TTS {OUTPUT_TTS}: {len(TTS)} paragraphs')


if __name__ == '__main__':
    main()
